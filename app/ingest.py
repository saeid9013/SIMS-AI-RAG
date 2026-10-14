"""Local parsers. Word locations are blocks, never fabricated pages."""
import hashlib
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path
from pypdf import PdfReader
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from .config import MAX_PAGES, MAX_TEXT

class ParseError(ValueError):
    pass

def normalize(text):
    text = text.replace('ي','ی').replace('ك','ک').replace('\u200c',' ')
    return re.sub(r'\s+', ' ', text).strip()

def chunks(text, size=1600, overlap=240):
    text = normalize(text)
    start = 0
    while start < len(text):
        end = min(start+size, len(text))
        if end < len(text):
            boundary = text.rfind(' ', start+size//2, end)
            if boundary > start:
                end = boundary
        value = text[start:end].strip()
        if value:
            yield value
        if end == len(text):
            break
        start = max(start+1, end-overlap)

def ocr_page(path, number):
    if not shutil.which('pdftoppm') or not shutil.which('tesseract'):
        raise ParseError('برای OCR نصب poppler و tesseract با زبان‌های fas و eng لازم است.')
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp)/'page'
        subprocess.run(['pdftoppm','-f',str(number),'-l',str(number),'-singlefile',
                        '-scale-to','2200','-png',str(path),str(out)], check=True,
                        capture_output=True, timeout=90)
        result = subprocess.run(['tesseract',str(out)+'.png','stdout','-l','fas+eng'],
                                check=True,capture_output=True,timeout=90)
        return result.stdout.decode('utf-8')

def extract(path, suffix, ocr=False):
    warnings, blocks, total = [], [], 0
    def add(location, text):
        nonlocal total
        total += len(text)
        if total > MAX_TEXT:
            raise ParseError('حجم متن استخراج‌شده از سقف مجاز بیشتر است.')
        blocks.append((location,text))
    if suffix == '.pdf':
        reader = PdfReader(path)
        if reader.is_encrypted:
            raise ParseError('PDF رمزگذاری شده است؛ نسخه بدون رمز بارگذاری کنید.')
        if len(reader.pages) > MAX_PAGES:
            raise ParseError('تعداد صفحات بیش از سقف تنظیم‌شده است.')
        for i, page in enumerate(reader.pages,1):
            text = page.extract_text() or ''
            label = f'صفحه {i}'
            if len(normalize(text)) < 35:
                if ocr:
                    text = ocr_page(path,i)
                    label += ' · OCR'
                    if len(normalize(text)) < 35:
                        warnings.append(f'صفحه {i}: حتی پس از OCR متن کافی ندارد؛ بازبینی لازم است.')
                else:
                    warnings.append(f'صفحه {i}: متن کافی ندارد؛ OCR را فعال و بازپردازش کنید.')
            add(label,text)
    elif suffix == '.docx':
        with zipfile.ZipFile(path) as archive:
            if sum(i.file_size for i in archive.infolist()) > 100*1024*1024:
                raise ParseError('حجم بازشده Word بیش از سقف مجاز است.')
        document = Document(path)
        heading = 'بدون عنوان'
        # Preserve paragraph/table order, unlike document.paragraphs + document.tables.
        for i, element in enumerate(document.element.body,1):
            if element.tag.endswith('}p'):
                p = Paragraph(element,document)
                if p.style and p.style.name.startswith('Heading'):
                    heading = p.text[:120]
                text = p.text
            elif element.tag.endswith('}tbl'):
                table = Table(element,document)
                text = '\n'.join(' | '.join(c.text for c in row.cells) for row in table.rows)
            else:
                continue
            add(f'{heading} · بلوک {i}', text)
        warnings.append('Word: سرصفحه، پاصفحه، کادر متن و تصاویر استخراج نمی‌شوند؛ جدول‌های پیچیده بازبینی شوند.')
    elif suffix == '.doc':
        executable = shutil.which('libreoffice') or shutil.which('soffice')
        if not executable:
            raise ParseError('برای DOC قدیمی LibreOffice نصب کنید یا فایل را به DOCX تبدیل کنید.')
        with tempfile.TemporaryDirectory() as tmp:
            subprocess.run([executable,'-env:UserInstallation=file://'+tmp+'/profile','--headless',
                            '--convert-to','docx','--outdir',tmp,str(path)],
                           check=True,capture_output=True,timeout=120)
            converted = Path(tmp)/(Path(path).stem+'.docx')
            if not converted.exists():
                raise ParseError('تبدیل DOC انجام نشد؛ فایل را دستی به DOCX تبدیل کنید.')
            return extract(converted,'.docx',ocr)
    elif suffix == '.txt':
        add('متن',Path(path).read_text(encoding='utf-8-sig'))
    else:
        raise ParseError('فرمت پشتیبانی نمی‌شود.')
    result = [(location,chunk) for location,text in blocks for chunk in chunks(text)]
    if not result:
        raise ParseError('هیچ متن قابل جست‌وجویی پیدا نشد. برای اسکن، OCR را فعال کنید.')
    return result, warnings
